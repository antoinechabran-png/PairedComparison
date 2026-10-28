"""
Topline engine for paired-comparison / sequential-monadic consumer tests.

Reads an INDEX sheet (data map) + DATA sheet, computes metrics per product
code, tests every candidate against a benchmark code and writes a colour-coded
Excel topline.

INDEX sheet layout expected (one column per variable, no header):
    row 0 : variable name (same as the DATA header)
    row 1 : question label
    row 2+: modalities written "code-label"  (e.g. "9-I like it extremely", "99#NA")
"""
from __future__ import annotations

import io
import math
import re
from dataclasses import dataclass, field

import numpy as np
import pandas as pd
from scipy import stats
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------- #
# Colours (sampled from the Givaudan threshold legend)
# --------------------------------------------------------------------------- #
COLORS = {
    ("pos", 99): "21A625", ("neg", 99): "A62521",
    ("pos", 95): "4BCB4F", ("neg", 95): "EB3933",
    ("pos", 90): "BBFF77", ("neg", 90): "FCE4D6",
    ("pos", 80): "DCE6F1", ("neg", 80): "FDE9D9",
}
WHITE_FONT_LEVELS = {("pos", 99), ("neg", 99), ("neg", 95)}

ALL_METRICS = ["Mean", "Std", "Top box", "Top 2 boxes", "Top 3 boxes",
               "Bottom box", "Bottom 2 boxes"]
BOX_DEF = {"Top box": ("top", 1), "Top 2 boxes": ("top", 2), "Top 3 boxes": ("top", 3),
           "Bottom box": ("bottom", 1), "Bottom 2 boxes": ("bottom", 2)}

MISSING_LABEL_RE = re.compile(
    r"^\s*(#?n/?a|na|n\.a\.|don'?t know|dk|refused?|not applicable|no answer|missing)\s*$",
    re.I)
MODALITY_RE = re.compile(r"^\s*(-?\d+(?:[.,]\d+)?)\s*[-#:=]\s*(.*)$")


# --------------------------------------------------------------------------- #
# Index parsing / data loading
# --------------------------------------------------------------------------- #
@dataclass
class Variable:
    name: str
    label: str
    codes: dict = field(default_factory=dict)   # code(float) -> label

    def code_label(self, code) -> str:
        try:
            c = float(code)
        except (TypeError, ValueError):
            return str(code)
        return self.codes.get(c, _fmt_code(c))

    def missing_codes(self) -> set:
        out = set()
        for c, lab in self.codes.items():
            if MISSING_LABEL_RE.match(str(lab)) or c >= 98:
                out.add(c)
        return out


def _fmt_code(c: float) -> str:
    return str(int(c)) if float(c).is_integer() else str(c)


def parse_index(raw: pd.DataFrame) -> dict[str, Variable]:
    """raw = INDEX sheet read with header=None."""
    out: dict[str, Variable] = {}
    for j in raw.columns:
        col = raw[j].tolist()
        if not col or pd.isna(col[0]):
            continue
        name = str(col[0]).strip()
        label = str(col[1]).strip() if len(col) > 1 and not pd.isna(col[1]) else name
        codes = {}
        for cell in col[2:]:
            if pd.isna(cell):
                continue
            m = MODALITY_RE.match(str(cell))
            if m:
                codes[float(m.group(1).replace(",", "."))] = m.group(2).strip() or m.group(1)
        out[name] = Variable(name, label, codes)
    return out


def load_workbook_sheets(file) -> dict[str, pd.DataFrame]:
    """Returns {sheet_name: raw dataframe with header=None}."""
    return pd.read_excel(file, sheet_name=None, header=None)


def build_dataset(raw_index: pd.DataFrame, raw_data: pd.DataFrame):
    """Link INDEX and DATA. Returns (data df, {data column -> Variable})."""
    index_vars = parse_index(raw_index)
    data = raw_data.copy()
    data.columns = [str(c).strip() for c in data.iloc[0]]
    data = data.iloc[1:].reset_index(drop=True)
    # de-duplicate header names if any
    seen, cols = {}, []
    for c in data.columns:
        if c in seen:
            seen[c] += 1
            cols.append(f"{c}.{seen[c]}")
        else:
            seen[c] = 0
            cols.append(c)
    data.columns = cols

    variables: dict[str, Variable] = {}
    idx_names = list(index_vars.keys())
    for pos, c in enumerate(data.columns):
        if c in index_vars:
            variables[c] = index_vars[c]
        elif pos < len(idx_names) and idx_names[pos] not in data.columns:
            # names differ but positions line up -> link by position
            v = index_vars[idx_names[pos]]
            variables[c] = Variable(c, v.label, dict(v.codes))
        else:
            variables[c] = Variable(c, c, {})
    for c in data.columns:
        conv = pd.to_numeric(data[c], errors="coerce")
        if conv.notna().sum() >= 0.8 * data[c].notna().sum() and data[c].notna().any():
            data[c] = conv
    return data, variables


def guess_roles(columns: list[str], variables: dict[str, Variable]) -> dict:
    up = {c: (c + " " + variables[c].label).upper() for c in columns}
    product = (next((c for c in columns if re.search(r"CODE", c.upper())), None)
               or next((c for c in columns if re.search(r"PRODUCT CODE|SAMPLE|PRODUIT|PRODUCT$", up[c])), None)
               or columns[min(1, len(columns) - 1)])
    cid = columns[0]
    qs = [c for c in columns if re.match(r"^Q", c, re.I) and c not in (product, cid)] or \
         [c for c in columns if c not in (product, cid)]
    pref = [c for c in qs if re.search(r"PREF", up[c])]
    liking = [c for c in qs if re.search(r"LIKING|OFR|APPRECIATION", up[c]) and c not in pref]
    strength = [c for c in qs if re.search(r"INTENSIT|INTENTISIT|STRENGTH", up[c])]
    used = set(liking) | set(strength) | set(pref) | {product, cid}
    frag = [c for c in qs if re.search(r"FRAGRANCE|SCENT|PARFUM", up[c]) and c not in used]
    used |= set(frag)
    other = [c for c in qs if c not in used]
    return dict(product=product, id=cid, liking=liking, strength=strength,
                fragrance=frag, other=other, preference=pref)


# --------------------------------------------------------------------------- #
# Statistics
# --------------------------------------------------------------------------- #
def _one_or_two_sided(stat_sign_diff: float, p_two: float, tail: str) -> float:
    """Convert a two-sided p (from a symmetric test) to the requested tail."""
    if tail == "two":
        return p_two
    if tail == "greater":
        return p_two / 2 if stat_sign_diff > 0 else 1 - p_two / 2
    return p_two / 2 if stat_sign_diff < 0 else 1 - p_two / 2


def mean_test(bench: pd.Series, cand: pd.Series, paired: bool, tail: str, equal_var: bool = True):
    """bench / cand: Series indexed by consumer id (valid values only)."""
    if paired:
        common = bench.index.intersection(cand.index)
        b, c = bench.loc[common].astype(float), cand.loc[common].astype(float)
        if len(common) < 3 or (c - b).std() == 0:
            return np.nan, "paired t"
        t, p = stats.ttest_rel(c, b)
        return _one_or_two_sided(t, p, tail), "paired t"
    b, c = bench.astype(float), cand.astype(float)
    if len(b) < 2 or len(c) < 2 or (b.std() == 0 and c.std() == 0):
        return np.nan, "t"
    t, p = stats.ttest_ind(c, b, equal_var=equal_var)
    return _one_or_two_sided(t, p, tail), ("Student t" if equal_var else "Welch t")


def z_test_props(x1: int, n1: int, x2: int, n2: int, tail: str):
    """Pooled two-proportion Z test, candidate (1) vs benchmark (2)."""
    if n1 == 0 or n2 == 0:
        return np.nan
    p1, p2 = x1 / n1, x2 / n2
    pp = (x1 + x2) / (n1 + n2)
    se = math.sqrt(pp * (1 - pp) * (1 / n1 + 1 / n2))
    if se == 0:
        return np.nan
    z = (p1 - p2) / se
    if tail == "two":
        return 2 * stats.norm.sf(abs(z))
    if tail == "greater":
        return stats.norm.sf(z)
    return stats.norm.cdf(z)


def mcnemar_test(bench_flag: pd.Series, cand_flag: pd.Series, tail: str):
    """McNemar on consumers who evaluated both products (bool Series by id).
    Exact binomial when discordant pairs < 25, else chi² with continuity correction."""
    common = bench_flag.index.intersection(cand_flag.index)
    b_ = cand_flag.loc[common].astype(bool) & ~bench_flag.loc[common].astype(bool)  # cand yes, bench no
    c_ = bench_flag.loc[common].astype(bool) & ~cand_flag.loc[common].astype(bool)
    b, c = int(b_.sum()), int(c_.sum())
    n = b + c
    if n == 0:
        return np.nan
    if n < 25:
        alt = {"two": "two-sided", "greater": "greater", "less": "less"}[tail]
        return stats.binomtest(b, n, 0.5, alternative=alt).pvalue
    chi2 = (abs(b - c) - 1) ** 2 / n
    p_two = stats.chi2.sf(chi2, 1)
    return _one_or_two_sided(b - c, p_two, tail)


def confidence_level(p: float, levels: list[int]) -> int | None:
    if p is None or (isinstance(p, float) and np.isnan(p)):
        return None
    conf = 1 - p
    for lv in sorted(levels, reverse=True):
        if conf >= lv / 100 - 1e-12:
            return lv
    return None


# --------------------------------------------------------------------------- #
# Topline computation
# --------------------------------------------------------------------------- #
@dataclass
class Settings:
    product_col: str
    id_col: str
    benchmark: float
    candidates: list
    sections: dict                      # section name -> [columns]
    metrics: list
    show_modalities: bool
    mean_tail: str = "two"              # two | greater | less
    prop_test: str = "z"                # z | mcnemar
    prop_tail: str = "two"
    equal_var: bool = True
    levels: tuple = (99, 95, 90)
    reversed_cols: list = field(default_factory=list)   # low code = best (for boxes)
    inverted_cols: list = field(default_factory=list)   # higher = worse (colour flipped)
    extra_missing: dict = field(default_factory=dict)   # col -> set of codes to drop
    preference: dict | None = None       # see compute_preference


@dataclass
class Cell:
    value: float
    kind: str              # mean | std | pct | n | text
    conf: int | None = None
    sign: int = 0          # +1 cand > bench, -1 cand < bench
    p: float | None = None


def is_paired(data: pd.DataFrame, s: Settings) -> bool:
    d = data[data[s.product_col].isin([s.benchmark] + list(s.candidates))]
    return d.groupby(s.id_col)[s.product_col].nunique().max() > 1


def compute_topline(data: pd.DataFrame, variables: dict[str, Variable], s: Settings):
    """Returns (rows, products, meta). Each row = dict(section, question, metric, cells{product: Cell})."""
    products = [s.benchmark] + [c for c in s.candidates if c != s.benchmark]
    paired = is_paired(data, s)
    prop_test = s.prop_test
    notes = []
    if prop_test == "mcnemar" and not paired:
        notes.append("McNemar requested but consumers did not evaluate several products "
                     "(independent samples) -> Z-test used instead.")
        prop_test = "z"

    sub = {p: data[data[s.product_col] == p] for p in products}
    rows = []
    rows.append(dict(section="", question="Base", metric="n",
                     cells={p: Cell(len(sub[p]), "n") for p in products}))

    for section, cols in s.sections.items():
        for col in cols:
            var = variables.get(col, Variable(col, col, {}))
            missing = var.missing_codes() | set(s.extra_missing.get(col, set()))
            per_prod = {}
            for p in products:
                ser = sub[p].set_index(s.id_col)[col].dropna()
                ser = ser[~ser.index.duplicated()]
                valid = ser[~ser.isin(missing)]
                per_prod[p] = (ser, valid)

            all_codes = sorted(set(var.codes) | set(pd.concat([v[0] for v in per_prod.values()]).unique()))
            scale = [c for c in all_codes if c not in missing]
            rev = col in s.reversed_cols
            flip = -1 if col in s.inverted_cols else 1
            qlabel = var.label
            first = True

            def add(metric, cells):
                nonlocal first
                rows.append(dict(section=section, question=qlabel if first else "",
                                 qname=col, metric=metric, cells=cells))
                first = False

            # ---- n answering
            add("Base (answered)", {p: Cell(len(per_prod[p][1]), "n") for p in products})

            # ---- mean / std
            if "Mean" in s.metrics and scale:
                cells = {}
                for p in products:
                    v = per_prod[p][1]
                    cell = Cell(v.mean() if len(v) else np.nan, "mean")
                    if p != s.benchmark:
                        pv, _ = mean_test(per_prod[s.benchmark][1], v, paired, s.mean_tail, s.equal_var)
                        _decorate(cell, pv, cells.get(s.benchmark), s.levels, flip)
                    cells[p] = cell
                add("Mean", cells)
            if "Std" in s.metrics and scale:
                add("Std", {p: Cell(per_prod[p][1].std(ddof=1) if len(per_prod[p][1]) > 1 else np.nan, "std")
                            for p in products})

            # ---- boxes
            ordered = sorted(scale, reverse=not rev)          # best first
            for metric in [m for m in s.metrics if m in BOX_DEF]:
                side, k = BOX_DEF[metric]
                if len(scale) < k + 1:
                    continue
                codes = ordered[:k] if side == "top" else ordered[-k:]
                lab = f"{metric} ({'+'.join(_fmt_code(c) for c in sorted(codes, reverse=True))})"
                add(lab, _prop_cells(per_prod, products, s, codes, valid_base=True,
                                     paired=paired, test=prop_test, flip=flip))

            # ---- % per modality
            if s.show_modalities:
                codes_out = sorted(all_codes, reverse=True) if scale else all_codes
                for c in codes_out:
                    add(f"{_fmt_code(c)} - {var.code_label(c)}",
                        _prop_cells(per_prod, products, s, [c], valid_base=False,
                                    paired=paired, test=prop_test, flip=flip))

    meta = dict(paired=paired, prop_test_used=prop_test, notes=notes,
                n_ids=data[s.id_col].nunique())
    return rows, products, meta


def _prop_cells(per_prod, products, s: Settings, codes, valid_base, paired, test, flip):
    cells = {}
    bench_ser, bench_valid = per_prod[s.benchmark]
    b_base = bench_valid if valid_base else bench_ser
    b_flag = b_base.isin(codes)
    for p in products:
        ser, valid = per_prod[p]
        base = valid if valid_base else ser
        flag = base.isin(codes)
        n, x = len(base), int(flag.sum())
        cell = Cell(x / n if n else np.nan, "pct")
        if p != s.benchmark and n and len(b_base):
            if test == "mcnemar":
                pv = mcnemar_test(b_flag, flag, s.prop_tail)
            else:
                pv = z_test_props(x, n, int(b_flag.sum()), len(b_base), s.prop_tail)
            _decorate(cell, pv, Cell(b_flag.mean(), "pct"), s.levels, flip, tail=s.prop_tail)
        cells[p] = cell
    return cells


def _decorate(cell: Cell, pv, bench_cell, levels, flip, tail=None):
    cell.p = pv
    lvl = confidence_level(pv, list(levels))
    if lvl is None or bench_cell is None or np.isnan(cell.value) or np.isnan(bench_cell.value):
        return
    diff = cell.value - bench_cell.value
    if diff == 0:
        return
    cell.conf = lvl
    cell.sign = (1 if diff > 0 else -1) * flip


# --------------------------------------------------------------------------- #
# Preference question
# --------------------------------------------------------------------------- #
PREF_TESTS = {
    "binomial_excl": "Binomial test vs 50% (no-preference excluded)",
    "binomial_split": "Binomial test vs 50% (no-preference split equally)",
    "z_dependent": "Z-test between the 2 preference proportions (dependent, total base)",
    "chi2": "Chi² goodness-of-fit (equal shares across all modalities)",
}


def compute_preference(data, variables, col, id_col, product_col, code_a, code_b,
                       code_np, test, levels, by_product=True, dedupe=True):
    """Returns list of groups: dict(name, n, pct{code: share}, p, winner)."""
    var = variables.get(col, Variable(col, col, {}))
    groups = [("Total", data)]
    if by_product:
        for code in sorted(data[product_col].dropna().unique()):
            groups.append((variables[product_col].code_label(code) if product_col in variables else code,
                           data[data[product_col] == code]))
    out = []
    for name, d in groups:
        ser = d[[id_col, col]].dropna()
        if dedupe and name == "Total":
            ser = ser.drop_duplicates(subset=id_col)
        ser = ser[col]
        n = len(ser)
        counts = ser.value_counts()
        codes = sorted(set(var.codes) | set(counts.index))
        pct = {c: (counts.get(c, 0) / n if n else np.nan) for c in codes}
        nA, nB = int(counts.get(code_a, 0)), int(counts.get(code_b, 0))
        nNP = int(counts.get(code_np, 0)) if code_np is not None else 0
        p = np.nan
        if n:
            if test == "binomial_excl" and nA + nB:
                p = stats.binomtest(nA, nA + nB, 0.5).pvalue
            elif test == "binomial_split":
                na = nA + nNP / 2
                nt = nA + nB + nNP
                if nt:
                    z = (na / nt - 0.5) / math.sqrt(0.25 / nt)
                    p = 2 * stats.norm.sf(abs(z))
            elif test == "z_dependent":
                pa, pb = nA / n, nB / n
                var_ = (pa + pb - (pa - pb) ** 2) / n
                if var_ > 0:
                    p = 2 * stats.norm.sf(abs((pa - pb) / math.sqrt(var_)))
            elif test == "chi2":
                obs = [counts.get(c, 0) for c in codes]
                if len(obs) > 1:
                    p = stats.chisquare(obs).pvalue
        lvl = confidence_level(p, list(levels))
        winner = 0
        if lvl and test != "chi2":
            winner = 1 if nA > nB else (-1 if nB > nA else 0)
        out.append(dict(name=str(name), n=n, pct=pct, p=p, conf=lvl, winner=winner,
                        label=var.label, codes=codes, var=var))
    return out


# --------------------------------------------------------------------------- #
# Excel writer
# --------------------------------------------------------------------------- #
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill("solid", fgColor="1F3864")
SEC_FILL = PatternFill("solid", fgColor="D9D9D9")
Q_FILL = PatternFill("solid", fgColor="F2F2F2")


def _fill_for(cell: Cell):
    if not cell.conf or not cell.sign:
        return None, False
    key = ("pos" if cell.sign > 0 else "neg", cell.conf)
    return PatternFill("solid", fgColor=COLORS[key]), key in WHITE_FONT_LEVELS


def write_excel(rows, products, meta, s: Settings, variables, pref_results=None,
                pref_info=None, title="Topline", include_pvalues=True) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Topline"
    pvar = variables.get(s.product_col)
    plabel = lambda p: pvar.code_label(p) if pvar else _fmt_code(p)

    tail_txt = {"two": "2-tailed", "greater": "1-tailed (candidate > benchmark)",
                "less": "1-tailed (candidate < benchmark)"}
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14, color="1F3864")
    ws["A2"] = (f"Benchmark: {plabel(s.benchmark)}  |  Means: "
                f"{'paired t-test' if meta['paired'] else ('Student t-test' if s.equal_var else 'Welch t-test')}"
                f" {tail_txt[s.mean_tail]}  |  Proportions: "
                f"{'McNemar' if meta['prop_test_used'] == 'mcnemar' else 'Z-test'} {tail_txt[s.prop_tail]}  |  "
                f"Design: {'paired (same consumers)' if meta['paired'] else 'independent samples'}")
    ws["A2"].font = Font(italic=True, size=9, color="595959")
    for i, note in enumerate(meta.get("notes", [])):
        ws.cell(row=3 + i, column=1, value="⚠ " + note).font = Font(size=9, color="C00000")

    r0 = 4 + len(meta.get("notes", []))
    headers = ["Section", "Question", "Metric"] + [
        f"{plabel(p)}{'  (BENCHMARK)' if p == s.benchmark else ''}" for p in products]
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=r0, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER

    pws = None
    if include_pvalues:
        pws = wb.create_sheet("Confidence levels")
        pws["A1"] = "Confidence of the gap vs benchmark (1 - p)"
        pws["A1"].font = Font(bold=True, size=12, color="1F3864")
        for j, h in enumerate(headers, 1):
            c = pws.cell(row=r0, column=j, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HDR_FILL
            c.alignment = Alignment(horizontal="center", wrap_text=True)

    r = r0 + 1
    last_section = None
    for row in rows:
        if row["section"] and row["section"] != last_section:
            ws.cell(row=r, column=1, value=row["section"].upper())
            for j in range(1, len(headers) + 1):
                ws.cell(row=r, column=j).fill = SEC_FILL
                ws.cell(row=r, column=j).font = Font(bold=True)
                ws.cell(row=r, column=j).border = BORDER
            last_section = row["section"]
            r += 1
        ws.cell(row=r, column=1, value="")
        qc = ws.cell(row=r, column=2, value=row["question"])
        if row["question"]:
            qc.font = Font(bold=True)
            for j in range(1, len(headers) + 1):
                ws.cell(row=r, column=j).fill = Q_FILL
        ws.cell(row=r, column=3, value=row["metric"])
        for j, p in enumerate(products, 4):
            cell = row["cells"][p]
            v = None if (cell.value is None or (isinstance(cell.value, float) and np.isnan(cell.value))) else float(cell.value)
            c = ws.cell(row=r, column=j, value=v)
            c.alignment = Alignment(horizontal="center")
            c.number_format = {"mean": "0.00", "std": "0.00", "pct": "0%", "n": "0"}.get(cell.kind, "General")
            fill, white = _fill_for(cell)
            if fill:
                c.fill = fill
                c.font = Font(bold=True, color="FFFFFF" if white else "000000")
            elif cell.kind == "n":
                c.font = Font(italic=True, color="595959")
            if pws is not None and cell.p is not None and not np.isnan(cell.p):
                pc = pws.cell(row=r, column=j, value=1 - cell.p)
                pc.number_format = "0.0%"
                if fill:
                    pc.fill = PatternFill("solid", fgColor=fill.fgColor.rgb)
        if pws is not None:
            for j in range(1, 4):
                pws.cell(row=r, column=j, value=ws.cell(row=r, column=j).value)
        for j in range(1, len(headers) + 1):
            ws.cell(row=r, column=j).border = BORDER
        r += 1

    for wsx in [ws] + ([pws] if pws is not None else []):
        wsx.column_dimensions["A"].width = 14
        wsx.column_dimensions["B"].width = 48
        wsx.column_dimensions["C"].width = 34
        for j in range(4, len(headers) + 1):
            wsx.column_dimensions[get_column_letter(j)].width = 16
        wsx.freeze_panes = wsx.cell(row=r0 + 1, column=4)
        wsx.row_dimensions[r0].height = 32
        wsx.page_setup.orientation = "landscape"
        wsx.page_setup.fitToWidth = 1
        wsx.page_setup.fitToHeight = 0
        wsx.sheet_properties.pageSetUpPr.fitToPage = True
        wsx.print_title_rows = f"{r0}:{r0}"

    # ---- preference sheet
    if pref_results:
        _write_pref(wb, pref_results, pref_info)

    # ---- legend
    _write_legend(wb, s, meta)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_pref(wb, results, info):
    ws = wb.create_sheet("Preference", 1)
    ws["A1"] = f"Preference – {results[0]['label']}"
    ws["A1"].font = Font(bold=True, size=13, color="1F3864")
    ws["A2"] = f"Test: {PREF_TESTS[info['test']]}  |  A = {info['a_label']}  vs  B = {info['b_label']}"
    ws["A2"].font = Font(italic=True, size=9, color="595959")
    hdr = ["Modality"] + [g["name"] for g in results]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center")
    ws.cell(row=5, column=1, value="Base (n)").font = Font(italic=True)
    for j, g in enumerate(results, 2):
        ws.cell(row=5, column=j, value=g["n"]).alignment = Alignment(horizontal="center")
    var = results[0]["var"]
    codes = results[0]["codes"]
    r = 6
    for code in codes:
        ws.cell(row=r, column=1, value=f"{_fmt_code(code)} - {var.code_label(code)}")
        for j, g in enumerate(results, 2):
            c = ws.cell(row=r, column=j, value=g["pct"].get(code, 0))
            c.number_format = "0%"
            c.alignment = Alignment(horizontal="center")
            if g["conf"] and g["winner"]:
                if (code == info["a"] and g["winner"] > 0) or (code == info["b"] and g["winner"] < 0):
                    key = ("pos", g["conf"])
                    c.fill = PatternFill("solid", fgColor=COLORS[key])
                    c.font = Font(bold=True, color="FFFFFF" if key in WHITE_FONT_LEVELS else "000000")
                elif (code == info["b"] and g["winner"] > 0) or (code == info["a"] and g["winner"] < 0):
                    key = ("neg", g["conf"])
                    c.fill = PatternFill("solid", fgColor=COLORS[key])
                    c.font = Font(bold=True, color="FFFFFF" if key in WHITE_FONT_LEVELS else "000000")
        r += 1
    ws.cell(row=r, column=1, value="Confidence (1-p)").font = Font(italic=True)
    for j, g in enumerate(results, 2):
        if g["p"] is not None and not np.isnan(g["p"]):
            c = ws.cell(row=r, column=j, value=1 - g["p"])
            c.number_format = "0.0%"
            c.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 40
    for j in range(2, len(hdr) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 14
    for row in ws.iter_rows(min_row=4, max_row=r, max_col=len(hdr)):
        for c in row:
            c.border = BORDER


def _write_legend(wb, s: Settings, meta):
    ws = wb.create_sheet("Legend")
    ws["A1"] = "Threshold"
    ws["B1"] = "Positive"
    ws["C1"] = "Negative"
    for c in ("A1", "B1", "C1"):
        ws[c].font = Font(bold=True)
        ws[c].border = BORDER
    for i, lv in enumerate(sorted(s.levels, reverse=True), 2):
        ws.cell(row=i, column=1, value=f"{lv}%").border = BORDER
        for col, side in ((2, "pos"), (3, "neg")):
            c = ws.cell(row=i, column=col)
            c.fill = PatternFill("solid", fgColor=COLORS[(side, lv)])
            c.border = BORDER
    r = len(s.levels) + 3
    lines = [
        "Candidates are compared with the benchmark. White = no significant gap.",
        "Positive = candidate above benchmark; Negative = candidate below "
        "(flipped for questions flagged 'higher = worse').",
        "Mean & Std: computed on valid scale answers (NA/don't know codes excluded).",
        "Box scores: % of valid scale answers. % per modality: % of all answers.",
        f"Design detected: {'paired – same consumers rated several products' if meta['paired'] else 'independent samples (monadic)'}.",
        f"Means test: {'paired t-test' if meta['paired'] else ('Student t-test (pooled variance)' if s.equal_var else 'Welch t-test')}.",
        f"Proportions test: {'McNemar (exact binomial when discordant pairs < 25)' if meta['prop_test_used'] == 'mcnemar' else 'pooled two-proportion Z-test'}.",
    ] + meta.get("notes", [])
    for i, t in enumerate(lines):
        ws.cell(row=r + i, column=1, value=t)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 18
